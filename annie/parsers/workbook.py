"""Excel workbooks as label tables (domain).

A label source may be an ``.xlsx`` workbook instead of a CSV. Each worksheet is read as
one table — the same ``header`` + ``rows`` shape :mod:`annie.parsers.csvmeta` produces for
a CSV — so everything downstream (key columns, value columns, joining) is format-agnostic.

Real workbooks are messier than CSVs, so reading is forgiving:

* the header is the **first non-empty row** (a sheet may start below or right of ``A1``);
* blank header cells are named after their column letter (``Column B``) and repeated
  names get a suffix (``id (2)``), so every column can be addressed;
* every cell becomes text with :func:`cell_text` — notably ``4021.0`` becomes ``"4021"``,
  because Excel stores ids as floats and they must still join to a video id.

Parsing is cached per file version, because the Dataset tab re-reads its sources often.
"""

from __future__ import annotations

import datetime as dt
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WORKBOOK_SUFFIXES = (".xlsx", ".xlsm")
"""File suffixes read as Excel workbooks (``.xls`` is a different, legacy format)."""

_Table = tuple[list[str], list[dict[str, str]]]
"""A parsed sheet: ``(header, rows)`` with every value as text."""


def is_workbook(path: str | Path) -> bool:
    """Return whether ``path`` names an Excel workbook (by suffix)."""
    return Path(path).suffix.lower() in WORKBOOK_SUFFIXES


def cell_text(value: object) -> str:
    """Return a cell value as the text Annie joins, tags and filters on.

    Args:
        value: A cell value as openpyxl returns it.

    Returns:
        ``""`` for an empty cell; integral numbers without a ``.0``; dates and times in
        ISO form (a date without a time part loses the ``00:00:00``); ``true``/``false``
        for booleans; anything else as stripped text.
    """
    if value is None:
        return ""
    if isinstance(value, bool):  # before int: bool is an int subclass
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, dt.datetime):
        if value.time() == dt.time():
            return value.date().isoformat()
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, dt.date | dt.time):
        return value.isoformat()
    return str(value).strip()


def sheet_names(path: str | Path) -> list[str]:
    """Return a workbook's sheet names in tab order, or ``[]`` if it can't be read."""
    file = Path(path)
    if not file.is_file():
        return []
    try:
        book = load_workbook(file, read_only=True)
    except (OSError, ValueError, KeyError):
        return []
    try:
        return list(book.sheetnames)
    finally:
        book.close()


def read_sheet(path: str | Path, sheet: str | None) -> _Table:
    """Read one worksheet as ``(header, rows)``, every value as text.

    Args:
        path: The workbook.
        sheet: The worksheet name; ``None`` reads the first sheet.

    Returns:
        The header (unique names) and one dict per non-empty data row, keyed by header.
        ``([], [])`` if the file or sheet is missing or unreadable.
    """
    file = Path(path)
    if not file.is_file():
        return [], []
    stat = file.stat()
    header, rows = _read_sheet_cached(str(file.resolve()), sheet, stat.st_mtime_ns, stat.st_size)
    # Copies, so a caller editing its rows can't change what the next caller reads.
    return list(header), [dict(row) for row in rows]


@lru_cache(maxsize=32)
def _read_sheet_cached(path: str, sheet: str | None, _mtime_ns: int, _size: int) -> _Table:
    """Parse a sheet; the file's mtime and size are part of the key so edits re-read."""
    try:
        book = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError, KeyError):
        return [], []
    try:
        name = sheet if sheet is not None else (book.sheetnames[0] if book.sheetnames else None)
        if name is None or name not in book.sheetnames:
            return [], []
        worksheet = book[name]
        worksheet.reset_dimensions()  # saved dimensions can be stale or missing; read all
        return _table(
            [[cell_text(v) for v in row] for row in worksheet.iter_rows(values_only=True)]
        )
    finally:
        book.close()


def _table(cells: list[list[str]]) -> _Table:
    """Turn text cells into ``(header, rows)``: first non-empty row is the header."""
    start = next((i for i, row in enumerate(cells) if any(row)), None)
    if start is None:
        return [], []
    raw = cells[start]
    width = max(i for i, value in enumerate(raw) if value) + 1
    header = _unique_names(raw[:width])
    rows: list[dict[str, str]] = []
    for row in cells[start + 1 :]:
        values = (row + [""] * width)[:width]
        if any(values):
            rows.append(dict(zip(header, values, strict=True)))
    return header, rows


def _unique_names(raw: list[str]) -> list[str]:
    """Name blank header cells by column letter and suffix repeats: ``id``, ``id (2)``."""
    names: list[str] = []
    seen: dict[str, int] = {}
    for index, value in enumerate(raw):
        base = value or f"Column {get_column_letter(index + 1)}"
        seen[base] = seen.get(base, 0) + 1
        names.append(base if seen[base] == 1 else f"{base} ({seen[base]})")
    return names
