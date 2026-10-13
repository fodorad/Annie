"""Tests for the Event-annotation domain: frame/second maths and JSON/CSV/XLSX export."""

from __future__ import annotations

import csv
import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from annie.dataset.events import (
    build_export_tree,
    clamp_event,
    effective_attributes,
    effective_color,
    events_to_csv_text,
    events_to_json_text,
    events_to_xlsx_bytes,
    export_events_csv,
    export_events_json,
    export_events_xlsx,
    frames_to_seconds,
    seconds_to_frame,
)
from annie.dataset.storage import EventRecord, EventType


def _event(
    video_id: str,
    track: str,
    start: int,
    end: int,
    *,
    label: str = "",
    note: str = "",
    attributes: dict[str, str] | None = None,
    color: str | None = None,
    event_type: str | None = None,
) -> EventRecord:
    return EventRecord(
        event_id=f"{video_id}-{track}-{start}",
        video_id=video_id,
        row_key=f"{video_id}::",
        track=track,
        start_frame=start,
        end_frame=end,
        label=label,
        note=note,
        attributes=attributes or {},
        color=color,
        event_type=event_type,
        updated_at="2020-01-01T00:00:00",
    )


class TestFrameSeconds(unittest.TestCase):
    def test_round_trip_at_common_fps(self) -> None:
        for fps in (25.0, 30.0, 23.976):
            for frame in (0, 1, 100, 4321):
                sec = frames_to_seconds(frame, fps)
                self.assertEqual(seconds_to_frame(sec, fps), frame)

    def test_zero_fps_degrades_safely(self) -> None:
        self.assertEqual(frames_to_seconds(100, 0.0), 0.0)
        self.assertEqual(seconds_to_frame(4.0, 0.0), 0)

    def test_seconds_to_frame_rounds_to_nearest(self) -> None:
        self.assertEqual(seconds_to_frame(0.041, 25.0), 1)  # 1.025 frames → 1
        self.assertEqual(seconds_to_frame(0.055, 25.0), 1)  # 1.375 frames → 1
        self.assertEqual(seconds_to_frame(0.061, 25.0), 2)  # 1.525 frames → 2
        self.assertEqual(seconds_to_frame(-1.0, 25.0), 0)  # never negative


class TestClampEvent(unittest.TestCase):
    def test_orders_reversed_pair(self) -> None:
        self.assertEqual(clamp_event(40, 10, 100), (10, 40))

    def test_bounds_to_last_frame(self) -> None:
        self.assertEqual(clamp_event(90, 200, 100), (90, 99))

    def test_clamps_negative_to_zero(self) -> None:
        self.assertEqual(clamp_event(-5, 10, 100), (0, 10))

    def test_unknown_frame_count_only_clamps_lower(self) -> None:
        # Before metadata loads, an event may still be authored past any upper bound.
        self.assertEqual(clamp_event(-3, 5000, 0), (0, 5000))


class TestExportTree(unittest.TestCase):
    def test_groups_by_video_and_track_with_seconds(self) -> None:
        events = [
            _event("v", "speech", 0, 25, label="A"),
            _event("v", "gesture", 50, 75),
        ]
        tree = build_export_tree(events, {"v": 25.0})
        self.assertEqual(tree["v"]["fps"], 25.0)
        tracks = tree["v"]["tracks"]
        self.assertEqual(set(tracks), {"speech", "gesture"})
        speech = tracks["speech"][0]
        self.assertEqual((speech["start_sec"], speech["end_sec"]), (0.0, 1.0))
        self.assertEqual(speech["label"], "A")

    def test_missing_fps_defaults_to_zero(self) -> None:
        tree = build_export_tree([_event("v", "speech", 0, 25)], {})
        self.assertEqual(tree["v"]["fps"], 0.0)
        self.assertEqual(tree["v"]["tracks"]["speech"][0]["start_sec"], 0.0)

    def test_participants_group_by_category_name(self) -> None:
        # The category name *is* the track, so exports already separate participants.
        events = [
            _event("v", "Mother", 0, 25, label="smile"),
            _event("v", "Baby", 50, 75, label="coo"),
        ]
        tree = build_export_tree(events, {"v": 25.0})
        self.assertEqual(set(tree["v"]["tracks"]), {"Mother", "Baby"})


class TestExportJson(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp())

    def test_json_nests_attributes(self) -> None:
        events = [_event("v", "speech", 0, 25, attributes={"speaker": "A"})]
        path = export_events_json(events, {"v": 25.0}, self.tmp / "out.json")
        data = json.loads(path.read_text(encoding="utf-8"))
        event = data["v"]["tracks"]["speech"][0]
        self.assertEqual(event["attributes"], {"speaker": "A"})

    def test_session_json_spans_multiple_videos(self) -> None:
        events = [_event("v", "speech", 0, 5), _event("w", "gesture", 0, 5)]
        path = export_events_json(events, {"v": 25.0, "w": 30.0}, self.tmp / "s.json")
        data = json.loads(path.read_text(encoding="utf-8"))
        self.assertEqual(set(data), {"v", "w"})


class TestExportCsv(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp())

    def _rows(self, path: Path) -> list[dict[str, str]]:
        with path.open(encoding="utf-8") as handle:
            return list(csv.DictReader(handle))

    def test_flattens_attributes_as_columns(self) -> None:
        events = [
            _event("v", "speech", 0, 25, attributes={"speaker": "A", "intensity": "high"}),
            _event("v", "speech", 30, 50, attributes={"speaker": "B"}),  # no intensity
        ]
        path = export_events_csv(events, {"v": 25.0}, self.tmp / "out.csv")
        rows = self._rows(path)
        self.assertEqual(rows[0]["attr_speaker"], "A")
        self.assertEqual(rows[0]["attr_intensity"], "high")
        # The second event lacks intensity → its column is present but empty (union header).
        self.assertEqual(rows[1]["attr_speaker"], "B")
        self.assertEqual(rows[1]["attr_intensity"], "")

    def test_fixed_columns_and_seconds(self) -> None:
        path = export_events_csv([_event("v", "speech", 0, 25)], {"v": 25.0}, self.tmp / "c.csv")
        row = self._rows(path)[0]
        self.assertEqual(row["video_id"], "v")
        self.assertEqual(row["track"], "speech")
        self.assertEqual(row["start_frame"], "0")
        self.assertEqual(row["end_sec"], "1.0")

    def test_no_attributes_yields_only_fixed_columns(self) -> None:
        path = export_events_csv([_event("v", "speech", 0, 5)], {"v": 25.0}, self.tmp / "n.csv")
        header = self._rows(path)[0].keys()
        self.assertNotIn("attr_", "".join(header))

    def test_color_column_present(self) -> None:
        events = [
            _event("v", "speech", 0, 5, color="#a8dadc"),
            _event("v", "speech", 10, 15),  # no colour → empty cell
        ]
        rows = self._rows(export_events_csv(events, {"v": 25.0}, self.tmp / "col.csv"))
        self.assertEqual(rows[0]["color"], "#a8dadc")
        self.assertEqual(rows[1]["color"], "")


class TestTextExport(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = Path(tempfile.mkdtemp())

    def test_json_text_matches_file(self) -> None:
        events = [_event("v", "speech", 0, 25)]
        path = export_events_json(events, {"v": 25.0}, self.tmp / "o.json")
        self.assertEqual(path.read_text(encoding="utf-8"), events_to_json_text(events, {"v": 25.0}))

    def test_csv_text_matches_file_and_has_header(self) -> None:
        events = [_event("v", "speech", 0, 25)]
        path = export_events_csv(events, {"v": 25.0}, self.tmp / "o.csv")
        text = events_to_csv_text(events, {"v": 25.0})
        self.assertEqual(path.read_bytes().decode("utf-8"), text)  # keeps CRLF as written
        self.assertTrue(text.startswith("video_id,track,start_frame"))


class TestMetaExport(unittest.TestCase):
    META = {"v": {"grp": "g1", "part": "a"}}

    def test_json_carries_meta_per_video(self) -> None:
        tree = json.loads(
            events_to_json_text(
                [_event("v", "speech", 0, 25), _event("w", "speech", 0, 5)], {}, self.META
            )
        )
        self.assertEqual(tree["v"]["meta"], {"grp": "g1", "part": "a"})
        self.assertNotIn("meta", tree["w"])

    def test_csv_has_meta_columns_between_fixed_and_attrs(self) -> None:
        text = events_to_csv_text(
            [_event("v", "speech", 0, 25), _event("w", "speech", 0, 5)], {}, self.META
        )
        header, first, second = text.splitlines()
        self.assertIn("color,meta_grp,meta_part", header)
        self.assertTrue(first.endswith(",g1,a"))
        self.assertTrue(second.endswith(",,"))  # video without meta → empty cells

    def test_no_meta_changes_nothing(self) -> None:
        events = [_event("v", "speech", 0, 25)]
        self.assertEqual(events_to_csv_text(events, {}), events_to_csv_text(events, {}, {}))
        self.assertNotIn("meta", build_export_tree(events, {})["v"])


_TYPES = {
    "Smile": EventType(
        name="Smile", ordinal=0, color="#b5e48c", attributes={"intensity": "1", "side": "l"}
    ),
    "Gaze away": EventType(name="Gaze away", ordinal=1),
}


class TestEffective(unittest.TestCase):
    """An event linked to an event category inherits its colour and attributes; its own win."""

    def test_colour_order(self) -> None:
        linked = _event("v", "Mother", 0, 5, event_type="Smile")
        self.assertEqual(effective_color(linked, _TYPES, "#ffffff"), "#b5e48c")
        own = _event("v", "Mother", 0, 5, event_type="Smile", color="#a2d2ff")
        self.assertEqual(effective_color(own, _TYPES, "#ffffff"), "#a2d2ff")
        no_colour = _event("v", "Mother", 0, 5, event_type="Gaze away")
        self.assertEqual(effective_color(no_colour, _TYPES, "#ffffff"), "#ffffff")
        gone = _event("v", "Mother", 0, 5, event_type="Deleted")
        self.assertIsNone(effective_color(gone, _TYPES))
        self.assertIsNone(effective_color(_event("v", "Mother", 0, 5)))

    def test_attributes_merge(self) -> None:
        event = _event("v", "Mother", 0, 5, event_type="Smile", attributes={"side": "r"})
        self.assertEqual(effective_attributes(event, _TYPES), {"intensity": "1", "side": "r"})
        plain = _event("v", "Mother", 0, 5, attributes={"a": "b"})
        self.assertEqual(effective_attributes(plain, _TYPES), {"a": "b"})
        self.assertEqual(effective_attributes(plain), {"a": "b"})


class TestExportWithEventTypes(unittest.TestCase):
    def setUp(self) -> None:
        self.events = [
            _event("v", "Mother", 0, 25, label="Smile", event_type="Smile"),
            _event("v", "Baby", 30, 50, label="look", color="#ffc8dd"),
        ]

    def test_json_has_type_and_effective_values(self) -> None:
        tree = json.loads(events_to_json_text(self.events, {"v": 25.0}, None, _TYPES))
        smile = tree["v"]["tracks"]["Mother"][0]
        self.assertEqual(smile["event_type"], "Smile")
        self.assertEqual(smile["color"], "#b5e48c")
        self.assertEqual(smile["attributes"], {"intensity": "1", "side": "l"})
        self.assertIsNone(tree["v"]["tracks"]["Baby"][0]["event_type"])

    def test_csv_has_type_and_effective_values(self) -> None:
        rows = list(csv.DictReader(io.StringIO(events_to_csv_text(self.events, {}, None, _TYPES))))
        self.assertEqual(rows[0]["event_type"], "Smile")
        self.assertEqual(rows[0]["color"], "#b5e48c")
        self.assertEqual(rows[0]["attr_intensity"], "1")
        self.assertEqual((rows[1]["event_type"], rows[1]["attr_intensity"]), ("", ""))


class TestExportXlsx(unittest.TestCase):
    """The Excel export: same columns as the CSV, typed numbers, a frozen and filtered header."""

    def _sheet(self, data: bytes):  # noqa: ANN202 - an openpyxl worksheet
        return load_workbook(io.BytesIO(data)).active

    def test_same_columns_as_csv_with_numbers(self) -> None:
        events = [
            _event("v", "Mother", 0, 25, label="Smile", event_type="Smile", note="ok"),
            _event("v", "Baby", 30, 50, attributes={"x": "y"}),
        ]
        meta = {"v": {"grp": "g1"}}
        sheet = self._sheet(events_to_xlsx_bytes(events, {"v": 25.0}, meta, _TYPES))
        header = [cell.value for cell in sheet[1]]
        csv_header = events_to_csv_text(events, {"v": 25.0}, meta, _TYPES).splitlines()[0]
        self.assertEqual(",".join(header), csv_header)
        first = dict(zip(header, [cell.value for cell in sheet[2]], strict=True))
        self.assertEqual(first["start_frame"], 0)
        self.assertEqual(first["end_sec"], 1.0)
        self.assertEqual((first["label"], first["event_type"]), ("Smile", "Smile"))
        self.assertEqual((first["meta_grp"], first["attr_intensity"]), ("g1", "1"))
        self.assertEqual(sheet.title, "Events")
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet.auto_filter.ref, f"A1:{sheet.cell(3, len(header)).coordinate}")
        self.assertTrue(sheet["A1"].font.bold)

    def test_long_text_caps_the_column_width(self) -> None:
        events = [_event("v", "Mother", 0, 5, note="n" * 500)]
        sheet = self._sheet(events_to_xlsx_bytes(events, {}))
        widths = [dim.width for dim in sheet.column_dimensions.values()]
        self.assertLessEqual(max(widths), 60)

    def test_no_events_is_a_header_only_sheet(self) -> None:
        sheet = self._sheet(events_to_xlsx_bytes([], {}))
        self.assertEqual(sheet.max_row, 1)
        self.assertEqual(sheet["A1"].value, "video_id")

    def test_writes_a_file(self) -> None:
        tmp = Path(tempfile.mkdtemp())
        path = export_events_xlsx([_event("v", "Mother", 0, 5)], {"v": 25.0}, tmp / "a" / "e.xlsx")
        self.assertEqual(self._sheet(path.read_bytes())["B2"].value, "Mother")


if __name__ == "__main__":
    unittest.main()
